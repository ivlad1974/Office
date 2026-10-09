import openpyxl, json, warnings
warnings.filterwarnings('ignore')

# --- Контрагенты from Реестр 2 (sheet named?) ---
wb2=openpyxl.load_workbook('/workspace/Реестр 2.xlsm',data_only=True)
print('R2 sheets:',wb2.sheetnames)

# --- Реестр ЗАКАЗОВ structure ---
wb3=openpyxl.load_workbook('/workspace/Реестр ЗАКАЗОВ 2.xlsm',data_only=True)
w=wb3['ПереченьЗаказов']
for r in range(1,10):
    vals=[w.cell(r,c).value for c in range(1,16)]
    print(r,[str(v)[:25] if v is not None else '' for v in vals])

# --- Контрагенты sheet inside Заказ workbook (values via data_only? links -> cached) ---
wb=openpyxl.load_workbook('/workspace/Заказ_3.1 1-Этап ДЗЕН от 11.08.2026.xlsm',data_only=True)
ko=wb['Контрагенты']
rows=[]
for r in range(2,78):
    row=[ko.cell(r,c).value for c in range(2,15)]
    if all(v in (None,0) for v in row): continue
    rows.append(row)
print('kontr rows:',len(rows))
json.dump({'kontr':rows},open('/workspace/kontr.json','w'),ensure_ascii=False,default=str)

# reestr headers full + sample with customer names
ws=wb2['РасчетДокументация']
hdr=[ws.cell(10,c).value for c in range(1,55)]
print('HDRS:',[ (i+1,h) for i,h in enumerate(hdr) if h])
