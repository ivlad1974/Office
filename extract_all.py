import openpyxl, json, re, warnings
warnings.filterwarnings('ignore')

def cellval(ws,r,c):
    v = ws.cell(row=r,column=c).value
    if v is None: return None
    if isinstance(v,str) and v.startswith('='): return None  # formula -> skip
    return v

out={}

# ---------- ЗАКАЗ ----------
wb=openpyxl.load_workbook('/workspace/Заказ_3.1 1-Этап ДЗЕН от 11.08.2026.xlsm',data_only=False)
m=wb['Материалы']
meta={
 'objectName': m.cell(1,5).value, 'org': m.cell(2,5).value, 'customer': m.cell(3,5).value,
 'version': m.cell(1,12).value, 'coefMat': m.cell(5,9).value or 1, 'coefWork': m.cell(5,10).value or 1,
}
rows=[]
for r in range(8,2003):
    b=m.cell(r,2).value; d=m.cell(r,4).value
    f=m.cell(r,6).value; g=m.cell(r,7).value; e=m.cell(r,5).value; k=m.cell(r,11).value
    if d is None and f is None: continue
    if isinstance(d,str) and d.startswith('='): d=None
    typ = b if isinstance(b,str) else None
    qty = f if isinstance(f,(int,float)) else None
    price = g if isinstance(g,(int,float)) else None
    rows.append({'t': typ, 'n': d if isinstance(d,str) else None, 'u': e if isinstance(e,str) else None,
                 'q': qty, 'p': price, 'note': k if isinstance(k,str) else None})
out['order_meta']=meta
out['order_rows']=rows

ks6=wb['КС-6']
k6={'stroyka':ks6.cell(7,3).value,'obekt':ks6.cell(9,3).value,'dogovorNum':ks6.cell(13,13).value,'dogovorDate':str(ks6.cell(14,13).value)[:10] if ks6.cell(14,13).value else None,
    'dates':[],'status':[]}
for c in range(6,15):
    v=ks6.cell(32,c).value
    k6['dates'].append(str(v)[:10] if v else None)
    st=ks6.cell(30,c).value
    k6['status'].append(st if isinstance(st,str) else '----')
out['ks6_header']=k6

# Контрагенты sheet inside заказ (link formula?)
try:
    ko=wb['Контрагенты']
    kontr=[]
    for r in range(3,8):
        row=[ko.cell(r,c).value for c in range(2,15)]
        row=[x if not(isinstance(x,str) and x.startswith('=')) else None for x in row]
        if any(row): kontr.append(row)
    out['our_companies']=kontr
except Exception as ex: out['our_companies']=[]

# ---------- РЕЕСТР 2 ----------
wb2=openpyxl.load_workbook('/workspace/Реестр 2.xlsm',data_only=True)
ws=wb2['РасчетДокументация']
hdr=[ws.cell(10,c).value for c in range(1,60)]
orders=[]
for r in range(11,900):
    vals=[ws.cell(r,c).value for c in range(1,60)]
    if all(v is None for v in vals): continue
    orders.append(vals)
out['reestr_headers']=hdr
out['reestr_orders']=orders[:1000]
print('reestr rows:',len(orders))

# Контрагенты в Реестре 2
for name in wb2.sheetnames:
    if 'онтраг' in name:
        wsk=wb2[name]
        our=[];cust=[]
        for r in range(2,80):
            row=[wsk.cell(r,c).value for c in range(2,15)]
            row=[v if not(isinstance(v,str) and v.startswith('=')) else None for v in row]
            if all(v in (None,0,'0') for v in row): continue
            if row[1] in (None,) and not row[2]: continue
            if r<8: our.append(row)
            else: cust.append(row)
        out['kontr_sheet_name']=name
        out['kontr_raw']=our+cust
        break

# ---------- РЕЕСТР ЗАКАЗОВ ----------
wb3=openpyxl.load_workbook('/workspace/Реестр ЗАКАЗОВ 2.xlsm',data_only=True)
print('rz sheets:',wb3.sheetnames)
wsz=wb3[wb3.sheetnames[0]]
for name in wb3.sheetnames:
    w=wb3[name]
    if w.max_row>50:
        wsz=w;break
hdr3=[wsz.cell(2,c).value for c in range(1,20)]
rz=[]
for r in range(3,1000):
    vals=[wsz.cell(r,c).value for c in range(1,20)]
    if all(v is None for v in vals): continue
    rz.append(vals)
out['rz_headers']=hdr3
out['rz_orders']=rz
print('rz rows:',len(rz))

json.dump(out,open('/workspace/webapp2/data/db.json','w'),ensure_ascii=False,default=str)
print('OK')
