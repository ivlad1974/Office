import openpyxl, json, warnings
warnings.filterwarnings('ignore')

# 1. КС-6 volumes (values) from заказ
wb=openpyxl.load_workbook('/workspace/Заказ_3.1 1-Этап ДЗЕН от 11.08.2026.xlsm',data_only=True)
ks6=wb['КС-6']
vols=[]
for r in range(34,2059):
    pos=ks6.cell(r,2).value  # B - № позиции
    if pos is None: continue
    row={'pos':pos}
    for i,c in enumerate(range(6,15)):  # F..N months
        v=ks6.cell(r,c).value
        if isinstance(v,(int,float)) and v!=0: row[f'm{i}']=v
    vols.append(row)
print('ks6 vols rows:',len(vols), vols[:3])

# 2. kontr structured
k=json.load(open('/workspace/kontr.json'))['kontr']
def mk(row, kind):
    def s(i): 
        v=row[i] if i<len(row) else None
        return '' if v in (None,0) else str(v)
    return {'kind':kind,'list':s(0),'name':s(1),'inn':s(2),'kpp':s(3),'okpo':s(4),'addr':s(5),
            'sign':s(6),'post':s(7),'fio':s(8),'nds':s(9),'header':s(10),'ks6str':s(12)}
our=[mk(r,'our') for r in k[1:4]]
cust_src=[r for r in k if r[0]=='Заказчики']
# customers: find section header then rows
idx=[i for i,r in enumerate(k) if r[0]=='Заказчики'][0]
custs=[]
for r in k[idx+1:]:
    m=mk(r,'customer')
    if m['list'] or m['name']: custs.append(m)
print('our:',our)
print('customers count:',len(custs), custs[:2])

# 3. Реестр ЗАКАЗОВ full columns
wb3=openpyxl.load_workbook('/workspace/Реестр ЗАКАЗОВ 2.xlsm',data_only=True)
w=wb3['ПереченьЗаказов']
hdr=[w.cell(10,c).value for c in range(1,22)]
rz=[]
for r in range(11,1000):
    vals=[w.cell(r,c).value for c in range(1,22)]
    if all(v in (None,0,'0','-') for v in vals): continue
    rz.append([str(v)[:200] if v is not None else '' for v in vals])
print('rz hdr:',hdr)
print('rz count:',len(rz),'sample:',rz[0])

json.dump({'ks6_vols':vols,'our':our,'customers':custs,'rz_hdr':[str(h) if h else '' for h in hdr],'rz':rz},
          open('/workspace/webapp2/data/db2.json','w'),ensure_ascii=False,default=str)
print('OK')
